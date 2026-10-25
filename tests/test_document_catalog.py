"""No-network dispatcher/cache-slot and actual native-format roundtrip tests."""
import asyncio
import io
import json
import os
from pathlib import Path
import struct
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import zipfile

_DATA = tempfile.TemporaryDirectory(prefix='ctxagent-catalog-')
os.environ.setdefault('CTXAGENT_DATA', _DATA.name)

from ctxagent import cache_policy, document_catalog as catalog, document_formats as formats
from ctxagent.agent import AgentSession, build_system_prompt
from ctxagent.config import CompactionConfig, SessionConfig
from ctxagent.segments import Conversation, K_SYSTEM, K_USER
from ctxagent.tool_batch import execute_batch, is_parallel_read


class CatalogTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.s = AgentSession('catalog-test', SessionConfig(workdirs=[self.tmp.name],
                             compaction=CompactionConfig(enabled=False)), Conversation())
        self.s.persist = lambda: None

    def open(self, name, data):
        p = self.root / name
        p.write_bytes(data.encode('utf-8') if isinstance(data, str) else data)
        doc = self.s.conv.workspace.open(self.s.sandbox(), str(p))
        version = next(v['version'] for v in self.s.conv.workspace.views if v['id'] == doc)
        return p, doc, version

    async def call(self, operation, **args):
        return await self.s._exec_tool('document_call',
                     {'operation': operation, 'arguments_json': json.dumps(args)})

    async def success(self, operation, **args):
        result = await self.call(operation, **args)
        self.assertFalse(result.error, result.text)
        return json.loads(result.text), result

    async def test_native_schema_stable_and_catalog_after_frozen_prefix(self):
        conv = self.s.conv
        conv.new('system', K_SYSTEM, build_system_prompt(self.s.cfg))
        conv.new('user', K_USER, 'fixed evidence')
        conv.freeze_prefix(2)
        conv.new('user', K_USER, 'working question')
        schema = self.s.tool_schema()
        self.assertIn('document_call', {t['function']['name'] for t in schema})
        self.assertEqual(catalog.messages(self.s), [])
        p, doc, version = self.open('data.json', '{"n":1}')
        self.assertEqual(schema, self.s.tool_schema())
        payload = catalog.messages(self.s)[0]['content']
        self.assertIn('document_json_update', payload)
        self.assertNotIn('document_pdf_render', payload)
        before = conv.to_messages()[:2] + catalog.messages(self.s) + conv.workspace.messages() + conv.to_messages()[2:]
        diagnostic = cache_policy.diagnose(before, schema, 'anthropic/test', 2, None, 'high')
        self.s.conv.workspace.close(doc)
        self.open('mesh.obj', 'v 0 0 0\nv 1 0 0\nv 0 1 0\nf 1 2 3\n')
        after = conv.to_messages()[:2] + catalog.messages(self.s) + conv.workspace.messages() + conv.to_messages()[2:]
        next_diag = cache_policy.diagnose(after, self.s.tool_schema(), 'anthropic/test', 2, diagnostic, 'high')
        self.assertEqual(before[:2], after[:2])
        self.assertEqual(schema, self.s.tool_schema())
        self.assertEqual(next_diag['shared_messages'], 2)
        self.assertNotIn('tool schema changed', next_diag['reasons'])
        self.assertIn('CATALOG', after[2]['content'])
        self.assertIn('LIVE DOCUMENT', after[3]['content'])
        self.assertGreater(self.s.stats()['document_catalog_tokens'], 0)

        captured = []
        async def stream(**kwargs):
            captured.append(kwargs)
            yield {'type': 'text', 'text': 'done'}
            yield {'type': 'done', 'finish_reason': 'stop'}
        with patch('ctxagent.agent.llm.stream_chat', side_effect=stream):
            async for event in self.s.run_turn(''):
                pass
        actual = captured[0]['messages']
        self.assertIn('CATALOG', actual[2]['content'])
        self.assertIn('LIVE DOCUMENT', actual[3]['content'])
        self.assertEqual(before[:2], actual[:2])

    async def test_dispatcher_validates_permissions_json_schema_and_closed_targets(self):
        p, doc, version = self.open('data.json', '{"n":1}')
        valid = {'document_id': doc, 'expected_version': version,
                 'updates': [{'pointer': '/n', 'value_json': '2'}]}
        for operation, raw in [('run_command', '{}'), ('document_json_update', '[]'),
                               ('document_json_update', '{"document_id":"a","document_id":"b"}'),
                               ('document_json_update', '{"x":NaN}'), ('document_json_update', '{}'),
                               ('document_json_update', json.dumps({**valid, 'ctx': {}}))]:
            result = await self.s._exec_tool('document_call', {'operation': operation, 'arguments_json': raw})
            self.assertTrue(result.error, result.text)
        data, _ = await self.success('document_json_update', **valid)
        self.assertEqual(json.loads(p.read_text()), {'n': 1})
        await self.success('document_apply_edits', plan_id=data['plan_id'])
        self.assertEqual(json.loads(p.read_text()), {'n': 2})
        self.s.cfg.enabled_tools.remove('document_json_update')
        result = await self.call('document_json_update', **valid)
        self.assertTrue(result.error)
        self.s.conv.workspace.close(doc)
        self.assertEqual(catalog.messages(self.s), [])
        result = await self.call('document_apply_edits', plan_id=data['plan_id'])
        self.assertTrue(result.error)

    def pdf_bytes(self):
        import fitz
        doc = fitz.open()
        for text in ('Page one', 'Page two'):
            doc.new_page().insert_text((72, 72), text)
        doc.set_toc([[1, 'First', 1]])
        raw = doc.tobytes()
        doc.close()
        return raw

    async def test_pdf_inspect_text_limits_render_and_scope(self):
        p, doc, version = self.open('data.pdf', self.pdf_bytes())
        args = {'document_id': doc, 'expected_version': version}
        data, _ = await self.success('document_pdf_inspect', **args)
        self.assertEqual(data['page_count'], 2)
        data, _ = await self.success('document_pdf_text', pages=[2], **args)
        self.assertIn('Page two', data['pages'][0]['text'])
        result = await self.call('document_pdf_text', pages=[1], max_chars=1, **args)
        self.assertTrue(result.error)
        data, rendered = await self.success('document_pdf_render', pages=[1], dpi=72, **args)
        self.assertEqual(len(rendered.images), 1)
        self.assertEqual(rendered.images[0]['source_sha256'], version)
        self.assertNotIn('images', data)
        result = await self.call('document_pdf_render', pages=[1], dpi=500, **args)
        self.assertTrue(result.error)
        with patch('ctxagent.document_formats.artifacts.render', return_value=[{'source_sha256': '0' * 64}]):
            result = await self.call('document_pdf_render', pages=[1], **args)
            self.assertTrue(result.error)
        _, other, other_version = self.open('other.json', '{"n":1}')
        result = await self.call('document_pdf_inspect', document_id=other, expected_version=other_version)
        self.assertTrue(result.error)
        p.write_bytes(self.pdf_bytes() + b'changed')
        result = await self.call('document_pdf_inspect', **args)
        self.assertTrue(result.error)

    async def test_readonly_dispatch_batches_overlap_but_edits_are_barriers(self):
        p, doc, version = self.open('data.pdf', self.pdf_bytes())
        active, peak = 0, 0
        lock = threading.Lock()
        overlap = threading.Event()
        def reader(*args, **kwargs):
            nonlocal active, peak
            with lock:
                active += 1
                peak = max(peak, active)
                if active >= 2:
                    overlap.set()
            overlap.wait(3)
            with lock:
                active -= 1
            return {'ok': True}
        calls = [{'id': str(i), 'function': {'name': 'document_call', 'arguments': json.dumps({
            'operation': 'document_pdf_inspect', 'arguments_json': json.dumps({
                'document_id': doc, 'expected_version': version})})}} for i in range(3)]
        results = []
        with patch('ctxagent.document_formats.pdf', side_effect=reader):
            async for phase, tc, name, args, result, dt in execute_batch(self.s, calls):
                if phase == 'result':
                    results.append(tc['id'])
                    self.assertFalse(result.error, result.text)
        self.assertGreaterEqual(peak, 2)
        self.assertEqual(results, ['0', '1', '2'])
        self.assertFalse(is_parallel_read('document_call', {'operation': 'document_xlsx_update'}))
        self.assertFalse(is_parallel_read('document_call', {'operation': 'document_apply_edits'}))

    def members(self, raw):
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            return {n: z.read(n) for n in z.namelist()}

    def with_member(self, raw, name, value):
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(out, 'w') as dest:
            for info in source.infolist():
                dest.writestr(info, source.read(info.filename))
            dest.writestr(name, value)
        return out.getvalue()

    def docx_bytes(self):
        from docx import Document
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('Hello ').bold = True
        p.add_run('world').italic = True
        out = io.BytesIO()
        doc.save(out)
        return self.with_member(out.getvalue(), 'customXml/opaque.bin', b'keep exactly\x00\xff')

    async def test_docx_review_run_styles_unrelated_members_and_exact_rollback(self):
        from docx import Document
        raw = self.docx_bytes()
        p, doc, version = self.open('data.docx', raw)
        args = {'document_id': doc, 'expected_version': version}
        info, _ = await self.success('document_docx_inspect', **args)
        self.assertEqual(info['paragraphs'][0]['text_runs'], 2)
        result = await self.call('document_docx_replace', updates=[
                                {'paragraph': 0, 'old_text': 'Hello world', 'new_text': 'changed'}], **args)
        self.assertTrue(result.error)
        plan, _ = await self.success('document_docx_replace', updates=[
                                     {'paragraph': 0, 'old_text': 'world', 'new_text': 'document'}], **args)
        self.assertEqual(p.read_bytes(), raw)
        self.assertIn('after_run', plan['change_summary'])
        await self.success('document_apply_edits', plan_id=plan['plan_id'])
        before, after = self.members(raw), self.members(p.read_bytes())
        self.assertEqual(set(before), set(after))
        self.assertEqual([n for n in before if before[n] != after[n]], ['word/document.xml'])
        actual = Document(str(p))
        self.assertEqual(actual.paragraphs[0].text, 'Hello document')
        self.assertTrue(actual.paragraphs[0].runs[0].bold)
        self.assertTrue(actual.paragraphs[0].runs[1].italic)
        await self.success('document_apply_edits', plan_id=plan['plan_id'], rollback=True)
        self.assertEqual(p.read_bytes(), raw)

    def xlsx_bytes(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        book = Workbook()
        sheet = book.active
        sheet.title = 'Data'
        sheet['A1'] = '001'
        sheet['A1'].font = Font(bold=True, color='00FF0000')
        sheet['B1'] = 3
        sheet['C1'] = '=B1*2'
        out = io.BytesIO()
        book.save(out)
        return self.with_member(out.getvalue(), 'custom/opaque.bin', b'preserve unknown part')

    async def test_xlsx_preserves_strings_styles_parts_and_refuses_formula_edits(self):
        from openpyxl import load_workbook
        raw = self.xlsx_bytes()
        p, doc, version = self.open('data.xlsx', raw)
        args = {'document_id': doc, 'expected_version': version}
        info, _ = await self.success('document_xlsx_inspect', **args)
        self.assertEqual(info['sheets'], ['Data'])
        for cell, value in [('C1', '4'), ('Z99', '4'), ('B1', '12345678901234567')]:
            result = await self.call('document_xlsx_update', sheet='Data', updates=[
                                     {'cell': cell, 'value_json': value}], **args)
            self.assertTrue(result.error, result.text)
        plan, _ = await self.success('document_xlsx_update', sheet='Data', updates=[
                                     {'cell': 'A1', 'value_json': '"002"'},
                                     {'cell': 'B1', 'value_json': '7'}], **args)
        self.assertEqual(p.read_bytes(), raw)
        await self.success('document_apply_edits', plan_id=plan['plan_id'])
        before, after = self.members(raw), self.members(p.read_bytes())
        self.assertEqual(set(before), set(after))
        self.assertEqual(set(n for n in before if before[n] != after[n]),
                         {'xl/worksheets/sheet1.xml', 'xl/workbook.xml'})
        book = load_workbook(str(p))
        self.assertEqual(book['Data']['A1'].value, '002')
        self.assertTrue(book['Data']['A1'].font.bold)
        self.assertEqual(book['Data']['B1'].value, 7)
        self.assertEqual(book['Data']['C1'].value, '=B1*2')
        self.assertTrue(book.calculation.fullCalcOnLoad)
        book.close()
        await self.success('document_apply_edits', plan_id=plan['plan_id'], rollback=True)
        self.assertEqual(p.read_bytes(), raw)

    async def test_signed_office_package_edit_refused_and_closed_plan_rejected(self):
        raw = self.with_member(self.docx_bytes(), '_xmlsignatures/signature.xml', b'<signature/>')
        p, doc, version = self.open('signed.docx', raw)
        result = await self.call('document_docx_replace', document_id=doc, expected_version=version,
                                updates=[{'paragraph': 0, 'old_text': 'world', 'new_text': 'document'}])
        self.assertTrue(result.error)
        self.assertIn('Signed', result.text)
        self.assertEqual(p.read_bytes(), raw)
        p, other, version = self.open('normal.docx', self.docx_bytes())
        plan, _ = await self.success('document_docx_replace', document_id=other, expected_version=version,
                                    updates=[{'paragraph': 0, 'old_text': 'world', 'new_text': 'document'}])
        self.s.conv.workspace.close(other)
        result = await self.call('document_apply_edits', plan_id=plan['plan_id'])
        self.assertTrue(result.error)

    async def test_obj_summary_transform_and_rollback_without_material_loading(self):
        raw = b'mtllib no_external_file.mtl\r\nv 0 0 0\r\nv 1 0 0\r\nv 0 1 0\r\nvn 0 0 1\r\nf 1//1 2//1 3//1\r\n'
        p, doc, version = self.open('data.obj', raw)
        self.assertIn('mesh geometry summary', self.s.conv.workspace.views[-1]['representation'])
        info, _ = await self.success('document_mesh_inspect', document_id=doc, expected_version=version)
        self.assertEqual(info['bounds']['max'], [1, 1, 0])
        plan, _ = await self.success('document_obj_transform', document_id=doc, expected_version=version,
                                     translate=[1, 2, 3], scale=2)
        self.assertEqual(p.read_bytes(), raw)
        self.assertEqual(plan['geometry_after']['bounds'], {'min': [1, 2, 3], 'max': [3, 4, 3]})
        await self.success('document_apply_edits', plan_id=plan['plan_id'])
        self.assertIn(b'vn 0 0 1\r\nf 1//1 2//1 3//1\r\n', p.read_bytes())
        await self.success('document_apply_edits', plan_id=plan['plan_id'], rollback=True)
        self.assertEqual(p.read_bytes(), raw)

    async def test_binary_stl_summary_and_invalid_geometry_refused(self):
        raw = bytes(80) + struct.pack('<I', 1) + struct.pack('<12fH', 0, 0, 1,
                                  0, 0, 0, 1, 0, 0, 0, 1, 0, 0)
        p, doc, version = self.open('data.stl', raw)
        info, _ = await self.success('document_mesh_inspect', document_id=doc, expected_version=version)
        self.assertEqual(info['faces'], 1)
        self.assertEqual(info['bounds']['max'], [1, 1, 0])
        self.assertNotIn('document_obj_transform', catalog.operation_names(self.s))
        for data, kind in [(b'v nan 0 0\n', '.obj'), (b'v 0 0 0\nf 1 2 3\n', '.obj'),
                           (raw[:-1], '.stl'), (b'solid ascii\nendsolid ascii', '.stl')]:
            with self.assertRaises(ValueError):
                formats.mesh_summary(data, kind)

    def test_xml_entities_and_oversized_package_are_refused(self):
        for encoding in ('utf-8', 'utf-16', 'utf-32'):
            with self.assertRaises(ValueError):
                formats._xml('<!DOCTYPE a [<!ENTITY b "payload">]><a>&b;</a>'.encode(encoding))
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr('oversized.xml', b'0' * (16 * 1024 * 1024 + 1))
        with self.assertRaises(ValueError):
            formats._package(output.getvalue())


if __name__ == '__main__':
    unittest.main(verbosity=2)
